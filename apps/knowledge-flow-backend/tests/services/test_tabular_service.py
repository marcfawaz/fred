from __future__ import annotations

import resource
from pathlib import Path
from types import SimpleNamespace

import pytest
from fred_core import KeycloakUser
from fred_core.common import OwnerFilter
from fred_core.documents.document_store import BaseDocumentMetadataStore as BaseMetadataStore
from fred_core.documents.document_structures import (
    DocumentMetadata,
    FileInfo,
    FileType,
    Identity,
    ProcessingStage,
    ProcessingStatus,
    SourceInfo,
    SourceType,
    Tagging,
)

from knowledge_flow_backend.application_context import ApplicationContext
from knowledge_flow_backend.core.processors.output.tabular_processor.tabular_processor import TabularProcessor
from knowledge_flow_backend.features.metadata.service import MetadataService
from knowledge_flow_backend.features.tabular.artifacts import (
    TABULAR_EXTENSION_KEY,
    document_artifact_prefix,
    read_tabular_artifact,
)
from knowledge_flow_backend.features.tabular.service import TabularDatasetAccessUnsupportedError, TabularService
from knowledge_flow_backend.features.tabular.structures import TabularQueryRequest
from knowledge_flow_backend.features.tag.structure import MissingTeamIdError


def _user() -> KeycloakUser:
    return KeycloakUser(
        uid="u-1",
        username="tester",
        email="tester@example.com",
        roles=["admin"],
    )


def _metadata(
    *,
    document_uid: str,
    file_name: str,
    tag_ids: list[str] | None = None,
    tag_names: list[str] | None = None,
) -> DocumentMetadata:
    return DocumentMetadata(
        identity=Identity(document_name=file_name, document_uid=document_uid, title=file_name),
        source=SourceInfo(source_type=SourceType.PUSH, source_tag="uploads"),
        file=FileInfo(file_type=FileType.CSV, mime_type="text/csv"),
        tags=Tagging(tag_ids=tag_ids or [], tag_names=tag_names or []),
    )


async def _ingest_csv(
    *,
    tmp_path: Path,
    metadata_store,
    document_uid: str,
    file_name: str,
    content: str,
    tag_ids: list[str] | None = None,
    tag_names: list[str] | None = None,
) -> DocumentMetadata:
    csv_path = tmp_path / file_name
    csv_path.write_text(content, encoding="utf-8")

    processor = TabularProcessor()
    metadata = _metadata(
        document_uid=document_uid,
        file_name=file_name,
        tag_ids=tag_ids,
        tag_names=tag_names,
    )
    processed_metadata = processor.process(str(csv_path), metadata)
    await MetadataService().save_document_metadata(_user(), processed_metadata)
    return processed_metadata


async def _ingest_excel_workbook(*, tmp_path: Path, document_uid: str) -> tuple[DocumentMetadata, Path]:
    """Run the spreadsheet input + output stages for a two-sheet workbook
    (Ventes: city/amount, Cibles: city/target) and save the metadata.

    Returns the saved metadata and the local output directory (output.md +
    tables.json), which callers can upload as document preview content.
    """
    from openpyxl import Workbook

    from knowledge_flow_backend.core.processors.input.excel_processor.excel_processor import ExcelProcessor
    from knowledge_flow_backend.core.processors.output.excel_processor.excel_table_registration_processor import ExcelTableRegistrationProcessor

    workbook_path = tmp_path / f"{document_uid}.xlsx"
    workbook = Workbook()
    ventes = workbook.active
    ventes.title = "Ventes"
    for row in [("city", "amount"), ("Paris", 10), ("Lyon", 20)]:
        ventes.append(row)
    cibles = workbook.create_sheet("Cibles")
    for row in [("city", "target"), ("Paris", 15), ("Lyon", 25)]:
        cibles.append(row)
    workbook.save(str(workbook_path))

    metadata = DocumentMetadata(
        identity=Identity(document_name=f"{document_uid}.xlsx", document_uid=document_uid, title=document_uid),
        source=SourceInfo(source_type=SourceType.PUSH, source_tag="uploads"),
        file=FileInfo(file_type=FileType.XLSX),
        tags=Tagging(tag_ids=[], tag_names=[]),
    )

    output_dir = tmp_path / f"{document_uid}-output"
    ExcelProcessor().convert_file_to_markdown(workbook_path, output_dir, document_uid)
    metadata = ExcelTableRegistrationProcessor().process(str(output_dir / "output.md"), metadata)
    await MetadataService().save_document_metadata(_user(), metadata)
    return metadata, output_dir


class _FakeRebac:
    def __init__(self, readable_document_uids: set[str]):
        self.readable_document_uids = readable_document_uids

    async def lookup_user_resources(self, user, permission):
        del user, permission
        return [SimpleNamespace(id=document_uid) for document_uid in sorted(self.readable_document_uids)]

    async def has_user_permission(self, user, permission, resource_id):
        del user, permission
        return resource_id in self.readable_document_uids


class _FakeTagService:
    """
    Minimal tag service stub used to emulate team/personal tabular scope.

    Why this exists:
    - Tabular tests need deterministic scope resolution without booting full tag
      ReBAC fixtures.

    How to use:
    - Configure readable tags globally and optional team-specific subsets.
    - Assign the instance to `service.tag_service`.
    """

    def __init__(
        self,
        *,
        readable_tag_ids: set[str],
        team_scopes: dict[str, set[str]] | None = None,
        personal_scope: set[str] | None = None,
    ) -> None:
        self.readable_tag_ids = readable_tag_ids
        self.team_scopes = team_scopes or {}
        self.personal_scope = personal_scope or set()

    async def list_authorized_tags_ids(self, user, owner_filter, team_id):
        del user
        if owner_filter is None:
            return set(self.readable_tag_ids)
        if owner_filter == OwnerFilter.TEAM:
            if not team_id:
                raise MissingTeamIdError("team_id is required when owner_filter is 'team'")
            return set(self.team_scopes.get(team_id, set()))
        return set(self.personal_scope)


class _TrackingMetadataStore(BaseMetadataStore):
    """
    Delegate wrapper that records targeted vs full metadata reads.

    Why this exists:
    - Tabular authorization should use one uid-targeted metadata lookup once
      ReBAC has already narrowed the readable document set.

    How to use:
    - Wrap the real metadata store, assign it to `service.metadata_store`, then
      inspect the recorded calls after one request.

    Example:
    - `service.metadata_store = _TrackingMetadataStore(metadata_store)`
    """

    def __init__(self, delegate) -> None:
        self._delegate = delegate
        self.get_all_metadata_calls = 0
        self.get_metadata_by_uids_calls: list[list[str]] = []

    async def get_all_metadata(self, filters: dict, session=None):
        self.get_all_metadata_calls += 1
        return await self._delegate.get_all_metadata(filters, session=session)

    async def get_metadata_by_uids(self, document_uids: list[str], session=None):
        self.get_metadata_by_uids_calls.append(list(document_uids))
        return await self._delegate.get_metadata_by_uids(document_uids, session=session)

    async def get_metadata_by_uid(self, document_uid: str, session=None) -> DocumentMetadata | None:
        return await self._delegate.get_metadata_by_uid(document_uid, session=session)

    async def get_metadata_in_tag(self, tag_id: str, session=None) -> list[DocumentMetadata]:
        return await self._delegate.get_metadata_in_tag(tag_id, session=session)

    async def list_by_source_tag(self, source_tag: str, session=None) -> list[DocumentMetadata]:
        return await self._delegate.list_by_source_tag(source_tag, session=session)

    async def save_metadata(self, metadata: DocumentMetadata, session=None) -> None:
        await self._delegate.save_metadata(metadata, session=session)

    async def delete_metadata(self, document_uid: str, session=None) -> None:
        await self._delegate.delete_metadata(document_uid, session=session)

    async def clear(self, session=None) -> None:
        await self._delegate.clear(session=session)

    def __getattr__(self, name):
        return getattr(self._delegate, name)


class _PresignedLocalContentStore:
    """
    Local content store wrapper that advertises presigned URLs.

    Why this exists:
    - Tests need a remote-style dataset location without requiring a live S3
      service.

    How to use:
    - Wrap the test local content store and override `get_presigned_url(...)`.
    """

    def __init__(self, delegate) -> None:
        self._delegate = delegate
        self.public_presigned_calls = 0
        self.internal_presigned_calls = 0

    def get_presigned_url(self, key, expires=None) -> str:
        del expires
        self.public_presigned_calls += 1
        return f"https://signed.example.invalid/{key}"

    def get_presigned_url_internal(self, key, expires=None) -> str:
        del expires
        self.internal_presigned_calls += 1
        return f"https://internal-signed.example.invalid/{key}"

    def __getattr__(self, name):
        return getattr(self._delegate, name)


class _UnsupportedRemoteContentStore:
    """
    Remote-style content store wrapper with no DuckDB-readable location support.

    Why this exists:
    - GCS can stream objects through the Python client while still lacking
      backend-internal signed URLs for DuckDB Parquet reads.

    How to use:
    - Wrap the local test content store to exercise the unsupported remote
      access path without a live object-storage service.
    """

    def __init__(self, delegate) -> None:
        self._delegate = delegate

    def get_presigned_url_internal(self, key, expires=None) -> str:
        del key, expires
        raise NotImplementedError("backend-internal signed URLs are unavailable")

    def __getattr__(self, name):
        return getattr(self._delegate, name)


@pytest.mark.asyncio
async def test_tabular_processor_stores_one_parquet_artifact_and_replaces_previous_revision(tmp_path, metadata_store):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    metadata = await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-1",
        file_name="sales.csv",
        content="city,amount,created_at\nParis,10,2024-01-01\nLyon,20,2024-01-02\n",
    )
    artifact = read_tabular_artifact(metadata)

    assert artifact is not None
    assert artifact.dataset_uid == "doc-1"
    assert artifact.row_count == 2
    assert [column.name for column in artifact.columns] == ["city", "amount", "created_at"]
    assert metadata.extensions is not None
    assert TABULAR_EXTENSION_KEY in metadata.extensions

    tabular_config = ApplicationContext.get_instance().get_config().storage.tabular_store
    object_prefix = document_artifact_prefix(
        artifacts_prefix=tabular_config.artifacts_prefix,
        document_uid="doc-1",
    )
    stored_objects = content_store.list_objects(object_prefix)
    assert len(stored_objects) == 1
    assert stored_objects[0].key == artifact.object_key

    updated_csv = tmp_path / "sales.csv"
    updated_csv.write_text("city,amount,created_at\nParis,30,2024-02-01\n", encoding="utf-8")

    processor = TabularProcessor()
    updated_metadata = processor.process(str(updated_csv), metadata)
    updated_artifact = read_tabular_artifact(updated_metadata)
    assert updated_artifact is not None
    assert updated_artifact.object_key != artifact.object_key

    stored_objects = content_store.list_objects(object_prefix)
    assert {stored_object.key for stored_object in stored_objects} == {
        artifact.object_key,
        updated_artifact.object_key,
    }

    await MetadataService().save_document_metadata(_user(), updated_metadata)

    stored_objects = content_store.list_objects(object_prefix)
    assert len(stored_objects) == 1
    assert stored_objects[0].key == updated_artifact.object_key


@pytest.mark.asyncio
async def test_tabular_processor_converts_csv_without_pandas_read_csv(tmp_path, monkeypatch):
    import pandas as pd

    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("city,amount\nParis,10\nLyon,20\n", encoding="utf-8")

    monkeypatch.setattr(pd, "read_csv", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("pandas path should not be used")))

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-no-pandas", file_name="sales.csv")

    processed_metadata = processor.process(str(csv_path), metadata)
    artifact = read_tabular_artifact(processed_metadata)

    assert artifact is not None
    assert artifact.row_count == 2
    assert [column.name for column in artifact.columns] == ["city", "amount"]
    assert [column.dtype for column in artifact.columns] == ["string", "integer"]
    assert processed_metadata.processing.stages[ProcessingStage.PREVIEW_READY] == ProcessingStatus.DONE
    assert processed_metadata.processing.stages[ProcessingStage.SQL_INDEXED] == ProcessingStatus.DONE


@pytest.mark.asyncio
async def test_tabular_processor_keeps_mixed_numeric_and_text_column_as_string(tmp_path):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    csv_path = tmp_path / "mixed-values.csv"
    with csv_path.open("w", encoding="utf-8") as file_handle:
        file_handle.write("id,label,value\n")
        for index in range(30_000):
            file_handle.write(f"{index},row,{index / 10}\n")
        file_handle.write("30000,tail,xxxxxxxxxxxxxxx\n")

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-mixed-types", file_name="mixed-values.csv")

    processed_metadata = processor.process(str(csv_path), metadata)
    artifact = read_tabular_artifact(processed_metadata)

    assert artifact is not None
    assert artifact.row_count == 30_001
    assert [column.name for column in artifact.columns] == ["id", "label", "value"]
    assert [column.dtype for column in artifact.columns] == ["integer", "string", "string"]


@pytest.mark.asyncio
async def test_tabular_processor_records_sample_values_for_low_cardinality_string_columns(tmp_path):
    """
    A low-cardinality string column carries its exact distinct values on the
    schema, so a SQL-writing agent sees the real stored casing instead of
    guessing it (e.g. it must not guess 'critical' when the data says
    'CRITICAL').
    """
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    rows = "\n".join(f"{i},{'CRITICAL' if i % 2 == 0 else 'LOW'}" for i in range(10))
    csv_path = tmp_path / "scan.csv"
    csv_path.write_text(f"id,severity\n{rows}\n", encoding="utf-8")

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-severity", file_name="scan.csv")
    processed_metadata = processor.process(str(csv_path), metadata)
    artifact = read_tabular_artifact(processed_metadata)

    assert artifact is not None
    severity_column = next(column for column in artifact.columns if column.name == "severity")
    assert severity_column.sample_values == ["CRITICAL", "LOW"]

    # Non-string columns are never sampled, regardless of cardinality.
    id_column = next(column for column in artifact.columns if column.name == "id")
    assert id_column.sample_values is None


@pytest.mark.asyncio
async def test_tabular_processor_skips_sample_values_for_high_cardinality_string_columns(tmp_path):
    """
    A string column above the low-cardinality threshold gets no sample_values
    — there is nothing useful to ground an agent's query with once every row
    is (close to) unique, and listing them all would bloat the schema payload.
    """
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    rows = "\n".join(f"{i},unique-label-{i}" for i in range(30))
    csv_path = tmp_path / "wide.csv"
    csv_path.write_text(f"row,label\n{rows}\n", encoding="utf-8")

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-high-cardinality", file_name="wide.csv")
    processed_metadata = processor.process(str(csv_path), metadata)
    artifact = read_tabular_artifact(processed_metadata)

    assert artifact is not None
    label_column = next(column for column in artifact.columns if column.name == "label")
    assert label_column.sample_values is None


@pytest.mark.asyncio
async def test_tabular_processor_cleans_temporary_parquet_after_upload(tmp_path, monkeypatch):
    """
    Ensure the DuckDB-generated Parquet temp file is deleted after content-store upload.

    Why this exists:
    - Tabular ingestion briefly materializes a Parquet artifact on local disk
      before uploading it to shared content storage.
    - The temporary file should not remain on disk once processing completes.

    How to use:
    - Patch the generated Parquet path to a known location, run
      `TabularProcessor.process(...)`, and assert the file existed during
      upload but not afterwards.
    """
    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("city,amount\nParis,10\nLyon,20\n", encoding="utf-8")

    tracked_parquet_path = tmp_path / "tracked-temp.parquet"
    observed: dict[str, object] = {"seen_during_upload": False}

    class _TrackedNamedTemporaryFile:
        def __enter__(self):
            self._handle = tracked_parquet_path.open("w+b")
            return self._handle

        def __exit__(self, exc_type, exc, tb):
            self._handle.close()

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-temp-cleanup", file_name="sales.csv")
    original_put_file = processor.content_store.put_file

    def _tracked_put_file(key: str, file_path: Path, *, content_type: str):
        observed["seen_during_upload"] = tracked_parquet_path.exists()
        assert file_path == tracked_parquet_path
        return original_put_file(key, file_path, content_type=content_type)

    monkeypatch.setattr(
        "knowledge_flow_backend.core.processors.output.tabular_processor.tabular_processor.tempfile.NamedTemporaryFile",
        lambda *args, **kwargs: _TrackedNamedTemporaryFile(),
    )
    monkeypatch.setattr(processor.content_store, "put_file", _tracked_put_file)

    processed_metadata = processor.process(str(csv_path), metadata)

    assert processed_metadata.file.row_count == 2
    assert observed["seen_during_upload"] is True
    assert not tracked_parquet_path.exists()


@pytest.mark.asyncio
async def test_tabular_processor_logs_parquet_artifact_write(tmp_path, caplog):
    """
    Ensure a successful Parquet upload leaves a durable, gathersable log line.

    Why this exists:
    - Before this fix, `_persist_parquet_artifact` uploaded the artifact to
      content storage without logging anything on success, so operators had
      no positive confirmation in stdout/OpenSearch that ingestion actually
      produced a queryable dataset.
    """
    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("city,amount\nParis,10\nLyon,20\n", encoding="utf-8")

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-log-check", file_name="sales.csv")

    with caplog.at_level("INFO"):
        processed_metadata = processor.process(str(csv_path), metadata)

    artifact = read_tabular_artifact(processed_metadata)
    assert artifact is not None

    tabular_log_records = [record.message for record in caplog.records if "[TABULAR]" in record.message]
    assert any("document_uid=doc-log-check" in message and f"object_key={artifact.object_key}" in message and "rows=2" in message for message in tabular_log_records)


@pytest.mark.integration
def test_tabular_processor_limits_python_rss_growth_for_large_csv(tmp_path):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    row_count = 400_000
    csv_path = tmp_path / "large.csv"
    with csv_path.open("w", encoding="utf-8") as file_handle:
        file_handle.write("city,amount,created_at\n")
        for index in range(row_count):
            file_handle.write(f"city_{index % 100},{index},2024-01-{(index % 28) + 1:02d}\n")

    assert csv_path.stat().st_size > 9 * 1024 * 1024

    processor = TabularProcessor()
    metadata = _metadata(document_uid="doc-large", file_name="large.csv")
    rss_before_kb = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss

    processed_metadata = processor.process(str(csv_path), metadata)

    rss_after_kb = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    rss_growth_bytes = max(0, rss_after_kb - rss_before_kb) * 1024

    assert processed_metadata.file.row_count == row_count
    assert rss_growth_bytes < 128 * 1024 * 1024


@pytest.mark.asyncio
async def test_tabular_service_lists_context_and_queries_datasets(tmp_path):
    app_context = ApplicationContext.get_instance()
    content_store = app_context.get_content_store()
    metadata_store = app_context.get_metadata_store()
    content_store.clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\nLyon,20\n",
    )
    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-targets",
        file_name="targets.csv",
        content="city,target\nParis,15\nLyon,25\n",
    )

    service = TabularService()
    datasets = await service.list_datasets(_user())
    assert {dataset.document_uid for dataset in datasets} == {"doc-sales", "doc-targets"}

    dataset_by_uid = {dataset.document_uid: dataset for dataset in datasets}
    sales_alias = dataset_by_uid["doc-sales"].query_alias
    targets_alias = dataset_by_uid["doc-targets"].query_alias

    sales_frame = await service.read_dataset_frame(_user(), "doc-sales")
    assert sales_frame.to_dict(orient="records") == [
        {"city": "Paris", "amount": 10},
        {"city": "Lyon", "amount": 20},
    ]

    query_response = await service.query_read(
        _user(),
        request=TabularQueryRequest(
            sql=(f"SELECT s.city, s.amount, t.target FROM {sales_alias} AS s JOIN {targets_alias} AS t ON s.city = t.city ORDER BY s.amount DESC"),
        ),
    )
    assert query_response.rows == [
        {"city": "Lyon", "amount": 20, "target": 25},
        {"city": "Paris", "amount": 10, "target": 15},
    ]

    count_query = await service.query_read(_user(), request=TabularQueryRequest(sql=f"SELECT COUNT(*) AS total_rows FROM {sales_alias}"))
    assert count_query.rows == [{"total_rows": 2}]

    with pytest.raises(ValueError, match="Only SELECT or WITH statements are allowed"):
        await service.query_read(
            _user(),
            request=TabularQueryRequest(sql=f"DROP TABLE {sales_alias}"),
        )


@pytest.mark.asyncio
async def test_tabular_service_rejects_duckdb_table_functions_outside_authorized_datasets(tmp_path):
    """
    Verify read-only tabular queries cannot bypass dataset scoping with DuckDB functions.

    Why this exists:
    - Bare table functions such as `read_parquet(...)` are not mounted Fred
      datasets and must be blocked by the dataset allowlist before execution.

    How to use:
    - Ingest one valid dataset, then assert that a query using
      `read_parquet(...)` is rejected.
    """

    app_context = ApplicationContext.get_instance()
    content_store = app_context.get_content_store()
    metadata_store = app_context.get_metadata_store()
    content_store.clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\nLyon,20\n",
    )

    service = TabularService()

    with pytest.raises(ValueError, match=r"unauthorized datasets: read_parquet\(\)"):
        await service.query_read(
            _user(),
            request=TabularQueryRequest(
                sql="SELECT * FROM read_parquet('/tmp/forbidden.parquet')",
            ),
        )


@pytest.mark.asyncio
async def test_tabular_service_rejects_explicit_dataset_requests_without_rebac_access(tmp_path, metadata_store):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    visible_metadata = await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-visible",
        file_name="visible.csv",
        content="city,amount\nParis,10\n",
    )
    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-hidden",
        file_name="hidden.csv",
        content="city,amount\nLyon,20\n",
    )

    service = TabularService()
    service.rebac = _FakeRebac({"doc-visible"})

    datasets = await service.list_datasets(_user())
    assert [dataset.document_uid for dataset in datasets] == [visible_metadata.document_uid]

    with pytest.raises(PermissionError, match="doc-hidden"):
        await service.query_read(
            _user(),
            request=TabularQueryRequest(
                sql="SELECT 1",
                dataset_uids=["doc-hidden"],
            ),
        )


@pytest.mark.asyncio
async def test_tabular_service_lists_authorized_datasets_with_targeted_metadata_lookup(tmp_path, metadata_store):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-visible",
        file_name="visible.csv",
        content="city,amount\nParis,10\n",
    )
    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-hidden",
        file_name="hidden.csv",
        content="city,amount\nLyon,20\n",
    )

    service = TabularService()
    tracking_store = _TrackingMetadataStore(metadata_store)
    service.metadata_store = tracking_store
    service.rebac = _FakeRebac({"doc-visible"})

    datasets = await service.list_datasets(_user())

    assert [dataset.document_uid for dataset in datasets] == ["doc-visible"]
    assert tracking_store.get_all_metadata_calls == 0
    assert tracking_store.get_metadata_by_uids_calls == [["doc-visible"]]


@pytest.mark.asyncio
async def test_tabular_service_requires_httpfs_for_remote_locations(tmp_path, metadata_store):
    """
    Verify remote tabular access fails clearly when `httpfs` cannot be loaded.

    Why this exists:
    - Remote S3-compatible access must stay `httpfs`-only with no hidden local
      copy fallback.

    How to use:
    - Wrap the local content store with presigned URLs and force
      `_ensure_httpfs_ready(...)` to fail.
    """

    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\nLyon,20\n",
    )

    service = TabularService()
    service.content_store = _PresignedLocalContentStore(content_store)
    service._ensure_httpfs_ready = lambda connection: (_ for _ in ()).throw(  # type: ignore[method-assign]
        RuntimeError("DuckDB httpfs is required for remote tabular dataset access.")
    )

    dataset = (await service.list_datasets(_user()))[0]
    with pytest.raises(RuntimeError, match="DuckDB httpfs is required"):
        await service.query_read(
            _user(),
            request=TabularQueryRequest(
                sql=f"SELECT city, amount FROM {dataset.query_alias} ORDER BY amount DESC",
            ),
        )

    with pytest.raises(RuntimeError, match="DuckDB httpfs is required"):
        await service.read_dataset_frame(_user(), "doc-sales")

    assert service.content_store.internal_presigned_calls >= 2
    assert service.content_store.public_presigned_calls == 0


@pytest.mark.asyncio
async def test_tabular_service_fails_cleanly_without_signed_url_or_local_path(tmp_path, metadata_store):
    """
    Verify unsupported content stores fail as an explicit tabular capability error.

    Why this exists:
    - GCS Workload Identity deployments can write Parquet artifacts before GCS
      V4 signed URLs are implemented, but DuckDB cannot read those artifacts
      through `get_object_stream(...)`.

    How to use:
    - The API layer maps this service exception to an unsupported-operation
      response instead of a generic 500.
    """

    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\nLyon,20\n",
    )

    service = TabularService()
    service.content_store = _UnsupportedRemoteContentStore(content_store)

    with pytest.raises(TabularDatasetAccessUnsupportedError, match="Unsupported operation"):
        await service.read_dataset_frame(_user(), "doc-sales")


@pytest.mark.asyncio
async def test_tabular_service_scopes_datasets_to_active_team_and_libraries(tmp_path, metadata_store):
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-team-a",
        file_name="sales-team-a.csv",
        content="city,amount\nParis,10\n",
        tag_ids=["tag-team-a"],
        tag_names=["Team A"],
    )
    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-team-b",
        file_name="sales-team-b.csv",
        content="city,amount\nLyon,20\n",
        tag_ids=["tag-team-b"],
        tag_names=["Team B"],
    )

    service = TabularService()
    service.tag_service = _FakeTagService(
        readable_tag_ids={"tag-team-a", "tag-team-b"},
        team_scopes={"team-a": {"tag-team-a"}, "team-b": {"tag-team-b"}},
    )

    team_a_datasets = await service.list_datasets(
        _user(),
        owner_filter=OwnerFilter.TEAM,
        team_id="team-a",
    )
    assert [dataset.document_uid for dataset in team_a_datasets] == ["doc-team-a"]

    [scoped_schema] = await service.describe_documents(
        _user(),
        ["doc-team-a"],
        owner_filter=OwnerFilter.TEAM,
        team_id="team-a",
    )
    assert scoped_schema.document_uid == "doc-team-a"
    assert scoped_schema.kind == "csv"

    assert (
        await service.list_datasets(
            _user(),
            owner_filter=OwnerFilter.TEAM,
            team_id="team-a",
            document_library_tags_ids=["tag-team-b"],
        )
    ) == []

    all_datasets = await service.list_datasets(_user())
    alias_by_uid = {dataset.document_uid: dataset.query_alias for dataset in all_datasets}

    scoped_rows = await service.query_read(
        _user(),
        request=TabularQueryRequest(
            sql=f"SELECT city, amount FROM {alias_by_uid['doc-team-a']}",
            owner_filter=OwnerFilter.TEAM,
            team_id="team-a",
        ),
    )
    assert scoped_rows.rows == [{"city": "Paris", "amount": 10}]

    with pytest.raises(ValueError, match="unauthorized datasets"):
        await service.query_read(
            _user(),
            request=TabularQueryRequest(
                sql=f"SELECT city, amount FROM {alias_by_uid['doc-team-b']}",
                owner_filter=OwnerFilter.TEAM,
                team_id="team-a",
            ),
        )


@pytest.mark.asyncio
async def test_tabular_service_expands_excel_workbook_into_multiple_datasets(tmp_path):
    """
    Vérifie la chaîne spreadsheet complète : input Excel (upload Parquet en clés
    canoniques + sidecar), output (promotion en `tabular_multi_v1`), puis
    éclatement 1 document → N datasets interrogeables en SQL avec les alias
    déterministes publiés dans le catalogue.
    """
    from knowledge_flow_backend.features.tabular.artifacts import build_table_query_alias, read_tabular_multi_artifact

    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    document_uid = "doc-excel"
    # Input stage (output.md + tables.json + Parquet en clés canoniques) puis
    # output stage (promotion du sidecar dans metadata.extensions)
    metadata, output_dir = await _ingest_excel_workbook(tmp_path=tmp_path, document_uid=document_uid)
    assert (output_dir / "output.md").is_file()
    multi = read_tabular_multi_artifact(metadata)
    assert multi is not None and len(multi.tables) == 2
    assert metadata.processing.stages[ProcessingStage.SQL_INDEXED] == ProcessingStatus.DONE

    # La sauvegarde (pruning des révisions) ne doit pas supprimer les tables courantes
    prefix = document_artifact_prefix(
        artifacts_prefix=ApplicationContext.get_instance().get_config().storage.tabular_store.artifacts_prefix,
        document_uid=document_uid,
    )
    assert {stored.key for stored in content_store.list_objects(prefix)} == {table.object_key for table in multi.tables}

    # Le service éclate le document en un dataset par table, alias stockés
    service = TabularService()
    datasets = await service.list_datasets(_user())
    ventes_alias = build_table_query_alias(document_uid, "Ventes", 1)
    cibles_alias = build_table_query_alias(document_uid, "Cibles", 1)
    assert {dataset.query_alias for dataset in datasets} == {ventes_alias, cibles_alias}
    assert {dataset.document_uid for dataset in datasets} == {document_uid}

    # Les deux tables se montent et se joignent en SQL
    joined = await service.query_read(
        _user(),
        request=TabularQueryRequest(
            sql=f"SELECT v.city, v.amount, c.target FROM {ventes_alias} AS v JOIN {cibles_alias} AS c ON v.city = c.city ORDER BY v.city",
        ),
    )
    assert [row["city"] for row in joined.rows] == ["Lyon", "Paris"]

    # Scoper la requête sur l'uid du document monte bien TOUTES ses tables
    scoped = await service.query_read(
        _user(),
        request=TabularQueryRequest(
            sql=f"SELECT COUNT(*) AS n FROM {cibles_alias}",
            dataset_uids=[document_uid],
        ),
    )
    assert scoped.rows == [{"n": 2}]


@pytest.mark.asyncio
async def test_tabular_service_lists_documents_mixing_csv_and_spreadsheet(tmp_path, metadata_store):
    """
    Vérifie la surface document-centrique : un document par entrée (CSV comme
    classeur), kind correct, tables imbriquées avec alias, et liste allégée
    (pas de colonnes — elles relèvent de l'endpoint schemas).
    """
    from knowledge_flow_backend.features.tabular.artifacts import build_table_query_alias

    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\nLyon,20\n",
    )
    await _ingest_excel_workbook(tmp_path=tmp_path, document_uid="doc-excel")

    service = TabularService()
    documents = await service.list_documents(_user())
    documents_by_uid = {document.document_uid: document for document in documents}
    assert set(documents_by_uid) == {"doc-sales", "doc-excel"}

    csv_document = documents_by_uid["doc-sales"]
    assert csv_document.kind == "csv"
    assert len(csv_document.tables) == 1
    assert csv_document.tables[0].sheet is None
    assert csv_document.tables[0].row_count == 2

    excel_document = documents_by_uid["doc-excel"]
    assert excel_document.kind == "spreadsheet"
    assert {table.sheet for table in excel_document.tables} == {"Ventes", "Cibles"}
    assert {table.query_alias for table in excel_document.tables} == {
        build_table_query_alias("doc-excel", "Ventes", 1),
        build_table_query_alias("doc-excel", "Cibles", 1),
    }

    # La liste reste légère : le résumé de table n'expose pas les colonnes.
    assert "columns" not in csv_document.tables[0].model_dump()


@pytest.mark.asyncio
async def test_tabular_service_describes_documents_in_batch_with_all_tables(tmp_path, metadata_store):
    """
    Vérifie que l'endpoint schemas est batch et table-complet : toutes les
    tables d'un classeur sont décrites (fin du « première table gagne »), et
    les uids interdits/inconnus lèvent 403/404 respectivement.
    """
    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\nLyon,20\n",
    )
    await _ingest_excel_workbook(tmp_path=tmp_path, document_uid="doc-excel")

    service = TabularService()
    schemas = await service.describe_documents(_user(), ["doc-sales", "doc-excel"])
    assert [schema.document_uid for schema in schemas] == ["doc-sales", "doc-excel"]

    csv_schema, excel_schema = schemas
    assert csv_schema.kind == "csv"
    assert [column.name for column in csv_schema.tables[0].columns] == ["city", "amount"]

    assert excel_schema.kind == "spreadsheet"
    assert len(excel_schema.tables) == 2
    columns_by_sheet = {table.sheet: [column.name for column in table.columns] for table in excel_schema.tables}
    assert columns_by_sheet == {"Ventes": ["city", "amount"], "Cibles": ["city", "target"]}

    with pytest.raises(ValueError, match="At least one document uid"):
        await service.describe_documents(_user(), [])

    service.rebac = _FakeRebac({"doc-sales", "doc-excel", "doc-unknown"})
    with pytest.raises(FileNotFoundError, match="doc-unknown"):
        await service.describe_documents(_user(), ["doc-sales", "doc-unknown"])

    service.rebac = _FakeRebac({"doc-sales"})
    with pytest.raises(PermissionError, match="doc-hidden"):
        await service.describe_documents(_user(), ["doc-sales", "doc-hidden"])


@pytest.mark.asyncio
async def test_tabular_service_document_markdown_is_spreadsheet_only(tmp_path, metadata_store):
    """
    Vérifie la route markdown : renvoie le catalogue `output.md` d'un classeur
    (avec les alias SQL exacts), refuse les documents non-spreadsheet (404) et
    les documents non lisibles (403).
    """
    from knowledge_flow_backend.features.tabular.artifacts import build_table_query_alias

    content_store = ApplicationContext.get_instance().get_content_store()
    content_store.clear()

    await _ingest_csv(
        tmp_path=tmp_path,
        metadata_store=metadata_store,
        document_uid="doc-sales",
        file_name="sales.csv",
        content="city,amount\nParis,10\nLyon,20\n",
    )
    metadata, output_dir = await _ingest_excel_workbook(tmp_path=tmp_path, document_uid="doc-excel")

    # Dans le pipeline réel, save_output uploade output.md et la preview passe
    # à DONE ; on rejoue ces deux étapes pour la lecture du catalogue.
    content_store.save_output("doc-excel", output_dir)
    metadata.mark_stage_done(ProcessingStage.PREVIEW_READY)
    await MetadataService().save_document_metadata(_user(), metadata)

    service = TabularService()
    content = await service.get_document_markdown(_user(), "doc-excel")
    assert "# Extraction summary" in content
    assert build_table_query_alias("doc-excel", "Ventes", 1) in content
    assert build_table_query_alias("doc-excel", "Cibles", 1) in content

    with pytest.raises(FileNotFoundError, match="not a spreadsheet"):
        await service.get_document_markdown(_user(), "doc-sales")

    service.rebac = _FakeRebac(set())
    with pytest.raises(PermissionError, match="doc-excel"):
        await service.get_document_markdown(_user(), "doc-excel")


def test_tabular_service_httpfs_install_is_attempted_after_load_failure():
    """
    Verify the tabular runtime tries `INSTALL httpfs` after a failed load.

    Why this exists:
    - Connected environments may start without the extension preloaded even
      though remote tabular access is still expected to work.

    How to use:
    - Call `_ensure_httpfs_ready(...)` with a fake DuckDB connection whose
      first `LOAD httpfs` fails and whose subsequent install+load succeeds.

    Example:
    - `service._ensure_httpfs_ready(fake_connection)`
    """

    class _ConnectionProbe:
        def __init__(self) -> None:
            self.commands: list[str] = []
            self._first_load = True

        def execute(self, sql: str) -> None:
            self.commands.append(sql)
            if sql == "LOAD httpfs" and self._first_load:
                self._first_load = False
                raise RuntimeError("missing extension")

    service = TabularService()
    connection = _ConnectionProbe()

    service._ensure_httpfs_ready(connection)  # type: ignore[arg-type]

    assert connection.commands == ["LOAD httpfs", "INSTALL httpfs", "LOAD httpfs"]


class _FakeVectorStore:
    """Records `add_documents` calls without touching a real vector backend."""

    def __init__(self, *, raise_on_add: bool = False):
        self.raise_on_add = raise_on_add
        self.added_documents: list = []

    def add_documents(self, documents):
        if self.raise_on_add:
            raise RuntimeError("simulated vector store failure")
        self.added_documents.extend(documents)
        return [doc.metadata.get("chunk_uid") for doc in documents]


def test_tabular_processor_pointer_chunks_disabled_by_default(tmp_path):
    """
    Verify the disabled-by-default path is untouched (RAG-DATASET-DISCOVERY-RFC.md §8:
    measured activation — no embedder/vector-store dependency paid unless enabled).
    """
    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("city,amount\nParis,10\nLyon,20\n", encoding="utf-8")

    processor = TabularProcessor()
    assert processor.tabular_config.pointer_chunks_enabled is False
    assert processor.vector_store is None
    assert processor.embedder is None

    metadata = _metadata(document_uid="doc-no-pointer", file_name="sales.csv")
    processed_metadata = processor.process(str(csv_path), metadata)

    artifact = read_tabular_artifact(processed_metadata)
    assert artifact is not None
    assert processed_metadata.processing.stages[ProcessingStage.SQL_INDEXED] == ProcessingStatus.DONE


def test_tabular_processor_emits_dataset_pointer_chunk_when_enabled(tmp_path):
    """
    Verify the pointer-chunk content, id, and kind match RAG-DATASET-DISCOVERY-RFC.md
    §2.1/§2.2/§2.4: one deterministic chunk, injection-resistant fixed template, no
    sample values in this increment.
    """
    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("city,amount\nParis,10\nLyon,20\n", encoding="utf-8")

    processor = TabularProcessor()
    processor.tabular_config.pointer_chunks_enabled = True
    fake_vector_store = _FakeVectorStore()
    processor.vector_store = fake_vector_store

    metadata = _metadata(document_uid="doc-pointer", file_name="sales.csv")
    processed_metadata = processor.process(str(csv_path), metadata)

    assert len(fake_vector_store.added_documents) == 1
    pointer_document = fake_vector_store.added_documents[0]

    assert pointer_document.metadata["chunk_uid"] == "doc-pointer::pointer"
    assert pointer_document.metadata["chunk_kind"] == "dataset_pointer"
    assert pointer_document.metadata["document_uid"] == "doc-pointer"
    # VectorSearchService unconditionally filters on metadata.retrievable=true
    # (vector_search_service.py:335,439,518) — without this the pointer chunk
    # is written but invisible to every search call. Caught live (2026-07-19)
    # by inspecting the actual OpenSearch document after a real ingestion.
    assert pointer_document.metadata["retrievable"] is True
    # Every deletion/consistency path in metadata/service.py gates its vector
    # cleanup on ProcessingStage.VECTORIZED being marked — without this, a
    # deleted tabular dataset leaves its pointer chunk permanently orphaned.
    # Caught live (2026-07-19): tabular artifacts/content were deleted but the
    # vector chunk was left behind because this guard never saw VECTORIZED.
    assert processed_metadata.processing.stages[ProcessingStage.VECTORIZED] == ProcessingStatus.DONE

    text = pointer_document.page_content
    assert "[DATASET POINTER" in text
    assert "[END DATASET POINTER]" in text
    assert "city" in text and "amount" in text
    assert "read_query" in text
    assert "dataset_uid=doc-pointer" in text
    # Increment 1 explicitly excludes sample values (§2.4) — guard against
    # accidentally reintroducing real cell values into the pointer text.
    assert "Paris" not in text
    assert "Lyon" not in text
    assert "Sample values" not in text


def test_tabular_processor_pointer_chunk_failure_is_non_fatal(tmp_path, caplog):
    """
    Verify a pointer-chunk write failure never breaks Parquet ingestion — the
    processor's primary contract (RAG-DATASET-DISCOVERY-RFC.md §2.1: best-effort).
    """
    csv_path = tmp_path / "sales.csv"
    csv_path.write_text("city,amount\nParis,10\nLyon,20\n", encoding="utf-8")

    processor = TabularProcessor()
    processor.tabular_config.pointer_chunks_enabled = True
    processor.vector_store = _FakeVectorStore(raise_on_add=True)

    metadata = _metadata(document_uid="doc-pointer-fail", file_name="sales.csv")
    with caplog.at_level("WARNING"):
        processed_metadata = processor.process(str(csv_path), metadata)

    artifact = read_tabular_artifact(processed_metadata)
    assert artifact is not None
    assert processed_metadata.processing.stages[ProcessingStage.SQL_INDEXED] == ProcessingStatus.DONE
    assert any("failed to write dataset pointer chunk" in record.message for record in caplog.records)


def test_tabular_store_config_pointer_chunks_disabled_by_default():
    """Guard the measured-activation default (RAG-DATASET-DISCOVERY-RFC.md §8)."""
    from knowledge_flow_backend.common.structures import TabularStoreConfig

    assert TabularStoreConfig().pointer_chunks_enabled is False


def test_allowed_chunk_keys_includes_chunk_kind():
    """chunk_kind must survive `sanitize_chunk_metadata`'s whitelist or pointer chunks silently lose it."""
    from knowledge_flow_backend.core.processors.output.vectorization_processor.vectorization_utils import sanitize_chunk_metadata

    clean, dropped = sanitize_chunk_metadata({"chunk_uid": "x::pointer", "chunk_kind": "dataset_pointer"})
    assert clean["chunk_kind"] == "dataset_pointer"
    assert "chunk_kind" not in dropped
